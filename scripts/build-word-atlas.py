import sys,json
from pathlib import Path
sys.path.insert(0,str(Path('.tmp/office-libs').resolve()))
from docx import Document
from docx.shared import Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path('docs/style-atlas');doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(25);sec.different_first_page_header_footer=True
for name,size,font,bold in [('Normal',12,'宋体',False),('Title',18,'黑体',True),('Heading 1',14,'黑体',True),('Heading 2',12,'宋体',True),('Heading 3',12,'宋体',True),('Caption',10.5,'宋体',False)]:
 st=doc.styles[name];st.font.name=font;st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor.from_string('24282D');st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font);st.paragraph_format.line_spacing=1.5;st.paragraph_format.space_before=Pt(12 if name.startswith('Heading') else 0);st.paragraph_format.space_after=Pt(6 if name.startswith('Heading') else 0);st.paragraph_format.first_line_indent=Pt(24 if name=='Normal' else 0)
def p(text,style='Normal',align=None):
 a=doc.add_paragraph(text,style);a.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if align is None else align;return a
def title(text):return p(text,'Heading 1',WD_ALIGN_PARAGRAPH.LEFT)
r=p('Office 文档基础样式范本','Title',WD_ALIGN_PARAGRAPH.CENTER);r.paragraph_format.space_before=Pt(120);r.paragraph_format.space_after=Pt(24);r.runs[0].font.size=Pt(28)
for s in ['一份文件 · 常用元素 · 可复用样式','封面 / 目录 / 章节 / 段落 / 表格 / 图片 / 页眉页脚','适用：一般企业办公文档','DEMO / 示例数据 · 2026年10月']:
 a=p(s,align=WD_ALIGN_PARAGRAPH.CENTER);a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.space_after=Pt(12)
doc.add_page_break();p('目录','Title',WD_ALIGN_PARAGRAPH.LEFT)
f=p('',align=WD_ALIGN_PARAGRAPH.LEFT);f.paragraph_format.first_line_indent=Pt(0)
for tag,text in [('begin',None),('instr',' TOC \\o "1-3" \\h \\z \\u '),('separate',None)]:
 e=OxmlElement('w:instrText' if tag=='instr' else 'w:fldChar');e.text=text if text else None
 if tag!='instr':e.set(qn('w:fldCharType'),tag)
 else:e.set(qn('xml:space'),'preserve')
 f.add_run()._r.append(e)
for text in ['一、章节、段落与列表\t3','二、表格、对齐与颜色\t4','三、图片与页眉页脚\t5']:p(text,align=WD_ALIGN_PARAGRAPH.LEFT)
e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),'end');doc.paragraphs[-1].add_run()._r.append(e)
doc.add_page_break();title('一、章节、段落与列表');p('（一）标题层级与正文','Heading 2',WD_ALIGN_PARAGRAPH.LEFT);p('1. 三级标题示例','Heading 3',WD_ALIGN_PARAGRAPH.LEFT)
p('正文使用宋体12pt、真实1.5倍行距、首行缩进24pt和两端对齐。标题通过Word全局样式定义，目录读取实际标题层级。按含义选择标题或正文样式，避免仅用加粗模拟章节。')
p('演示数据统一为预算300万元、实际支出290万元、余额10万元、执行率96.7%。数值只用于说明样式，正式使用时应替换为真实资料并核验。')
p('（二）按用途选择对齐方式','Heading 2',WD_ALIGN_PARAGRAPH.LEFT)
for text,align in [('左对齐：短说明、标签和列表。',0),('居中：独立标题与图注。',1),('右对齐：日期与签署位置。',2),('两端对齐：连续正文保持左右边界整齐，末行自然结束，不用手工空格凑齐。',3)]:
 a=p(text,align=align);a.paragraph_format.first_line_indent=Pt(0)
p('（三）真正的项目符号和编号','Heading 2',WD_ALIGN_PARAGRAPH.LEFT)
for text in ['并列事项使用项目符号，保持同级缩进。','状态同时使用文字说明，颜色仅作辅助。']:p(text,'List Bullet',WD_ALIGN_PARAGRAPH.LEFT)
for text in ['先确认内容与数据口径。','再应用元素样式并检查效果。']:p(text,'List Number',WD_ALIGN_PARAGRAPH.LEFT)
doc.add_page_break();title('二、表格、对齐与颜色');p('（一）真实表格与汇总','Heading 2',WD_ALIGN_PARAGRAPH.LEFT);p('表头加粗并跨页重复，金额右对齐，文字左对齐；单位与示例性质明确，不用整张表格图片代替可编辑数据。')
rows=[['类别','预算/万元','实际/万元','余额/万元','说明'],['研发','120','114','6','按计划'],['运营','80','84','-4','待复核'],['营销','60','54','6','按计划'],['支持','40','38','2','按计划'],['合计','300','290','10','96.7%']]
t=doc.add_table(rows=6,cols=5);t.style='Table Grid';t.autofit=False
for i,row in enumerate(rows):
 for j,text in enumerate(row):
  cell=t.cell(i,j);cell.text=text;a=cell.paragraphs[0];a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.line_spacing=1;a.alignment=WD_ALIGN_PARAGRAPH.RIGHT if j in [1,2,3] else WD_ALIGN_PARAGRAPH.LEFT
  for run in a.runs:run.font.size=Pt(11);run.font.bold=i in [0,5]
  if i in [0,5]:shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFEEEA');cell._tc.get_or_add_tcPr().append(shade)
rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
p('表1　预算与实际对比（DEMO）','Caption',WD_ALIGN_PARAGRAPH.CENTER);p('（二）强调、颜色与签署','Heading 2',WD_ALIGN_PARAGRAPH.LEFT)
a=p('默认正文以深色传递信息。重要判断使用加粗；辅助标签可少量使用哑铜色，不能只靠颜色区分状态。');a.runs[0].bold=True
a=p('辅助标签：统计口径 / 示例数据 / 待复核',align=WD_ALIGN_PARAGRAPH.LEFT);a.paragraph_format.first_line_indent=Pt(0);a.runs[0].font.color.rgb=RGBColor.from_string('836344');p('复核人：文档编制者与复核者',align=WD_ALIGN_PARAGRAPH.RIGHT)
doc.add_page_break();title('三、图片与页眉页脚');p('（一）嵌入图片、图注与替代文字','Heading 2',WD_ALIGN_PARAGRAPH.LEFT);p('图片保持原始宽高比，并使用替代文字和独立图注说明用途。下图为项目原创PNG，作为真实图片嵌入文件。')
pic=doc.add_picture(str(root/'assets/original-workflow.png'),width=Mm(155));pic._inline.docPr.set('descr','规划、制作、复核三个阶段的原创流程示意图');p('图1　从内容规划到效果复核（原创示意图）','Caption',WD_ALIGN_PARAGRAPH.CENTER)
p('（二）页眉、页脚与真实字段','Heading 2',WD_ALIGN_PARAGRAPH.LEFT);p('封面不显示页眉页脚。后续页面统一显示页眉与居中动态PAGE页码。目录为真实TOC字段，修改标题后应更新。范本中的字段不表示桥接已具备所有字段重建能力。')
p('（三）按元素复用','Heading 2',WD_ALIGN_PARAGRAPH.LEFT);p('先另存工作副本，再替换演示内容。保留标题层级、表格结构与图片比例；完成后更新目录，并逐页检查标题孤行、表格跨页、图片边界和页码。')
a=sec.header.paragraphs[0];a.text='Office 基础样式范本 · DEMO';a.paragraph_format.first_line_indent=Pt(0)
a=sec.footer.paragraphs[0];a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');a._p.append(field)
for area in [sec.header,sec.footer]:
 for a in area.paragraphs:
  for r in a.runs:r.font.name='宋体';r.font.size=Pt(10.5)
doc.save(root/'word-base-styles.docx');print('Word OOXML atlas saved: semantic styles, real TOC/PAGE, table and image')
